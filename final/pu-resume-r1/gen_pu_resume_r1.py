# -*- coding: utf-8 -*-
"""靜宜 E-Portfolio 履歷自傳 r1。只填已核過的事實；無依據欄位留空。"""
import os, re, hashlib, shutil
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from weasyprint import HTML
import pymupdf
from PIL import Image

ROOT = "/tmp/pu-resume-r1"
OUT_DIR = "/opt/cursor/artifacts/pu-resume-r1"
STORE = "/cursor/stores/self/final/pu-resume-r1"
WS_DIR = "/workspace/final/pu-resume-r1"
PHOTO = os.path.join(ROOT, "photo.jpg")
FONT_SANS = "/tmp/fju-r19/fonts/NotoSansTC-VF.ttf"
FONT_SERIF = "/tmp/fju-r19/fonts/NotoSerifTC-VF.ttf"
for d in (OUT_DIR, STORE, WS_DIR, os.path.join(OUT_DIR, "pages"), os.path.join(STORE, "pages"), os.path.join(WS_DIR, "pages")):
    os.makedirs(d, exist_ok=True)

DOCX_NAME = "履歷自傳_顧晉瑋.docx"
PDF_NAME = "履歷自傳_顧晉瑋.pdf"

# --- facts ---
NAME = "顧晉瑋"
GENDER = "男"
BIRTH = "民國 94 年 7 月 29 日"
PHONE = "0966198826"
EMAIL = "demohan513@gmail.com"
ADDR = "402 台中市南區學府路200號14樓之1"
SID = "411211325"
UNI = "靜宜大學"
DEPT = "資訊管理學系"
HS = "楊子高中"
EN_NAME = "KU CHIN-WEI"


def set_run_font(run, size=11, bold=False, name="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), "標楷體")
    rFonts.set(qn("w:cs"), name)


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), kwargs.get("val", "single"))
        el.set(qn("w:sz"), kwargs.get("sz", "8"))
        el.set(qn("w:color"), kwargs.get("color", "1A4550"))
        el.set(qn("w:space"), "0")


def vmerge(cell, val):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    v = OxmlElement("w:vMerge")
    if val:
        v.set(qn("w:val"), val)
    tcPr.append(v)


def set_widths(table, cms):
    table.autofit = False
    table.allow_autofit = False
    total = sum(cms)
    table.w = None
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(cms[i])
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(cms[i] * 567)))
            tcW.set(qn("w:type"), "dxa")


def cell_text(cell, text, size=10.5, bold=False, align="left", fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    # vertical align
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)
    return p


def add_heading_p(doc, text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True)
    return p


def add_sub(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A4550")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.45
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    # default style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "標楷體")

    k = doc.add_paragraph()
    k.alignment = WD_ALIGN_PARAGRAPH.CENTER
    k.paragraph_format.space_after = Pt(0)
    rk = k.add_run("靜宜大學 E-Portfolio")
    set_run_font(rk, size=10)
    add_heading_p(doc, "履歷表")

    t = doc.add_table(rows=7, cols=5)
    t.style = "Table Grid"
    cell_text(t.cell(0, 0), "應徵職務", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(0, 1), "", size=10.5)
    t.cell(0, 1).merge(t.cell(0, 3))

    cell_text(t.cell(1, 0), "姓名", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(1, 1), NAME, size=10.5)
    cell_text(t.cell(1, 2), "性別", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(1, 3), GENDER, size=10.5)

    cell_text(t.cell(2, 0), "出生年月日", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(2, 1), BIRTH, size=10.5)
    t.cell(2, 1).merge(t.cell(2, 3))

    cell_text(t.cell(3, 0), "服役狀況", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(3, 1), "□免役　　□未役　　□待役　　□役畢　　年　　月", size=9)
    t.cell(3, 1).merge(t.cell(3, 3))

    cell_text(t.cell(4, 0), "聯絡電話", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(4, 1), PHONE, size=10.5)
    t.cell(4, 1).merge(t.cell(4, 3))

    cell_text(t.cell(5, 0), "電子信箱", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(5, 1), EMAIL, size=10.5)
    t.cell(5, 1).merge(t.cell(5, 3))

    cell_text(t.cell(6, 0), "聯絡地址", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(t.cell(6, 1), ADDR, size=10.5)
    t.cell(6, 1).merge(t.cell(6, 3))

    photo_cell = t.cell(0, 4)
    photo_cell.merge(t.cell(6, 4))
    photo_cell.text = ""
    p = photo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = p.add_run("請貼照片\n")
    set_run_font(cap, size=8)
    run = p.add_run()
    run.add_picture(PHOTO, width=Cm(2.5))
    set_cell_border(photo_cell)
    shade(photo_cell, "FFFFFF")
    set_widths(t, [3.0, 4.0, 2.4, 4.2, 3.8])

    add_sub(doc, "教育程度")
    edu = doc.add_table(rows=3, cols=3)
    edu.style = "Table Grid"
    cell_text(edu.cell(0, 0), "學校", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(edu.cell(0, 1), "科系所", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(edu.cell(0, 2), "修業年限", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(edu.cell(1, 0), HS, size=10.5)
    cell_text(edu.cell(1, 1), "", size=10.5)  # 科系 unknown
    cell_text(edu.cell(1, 2), "", size=10.5)
    cell_text(edu.cell(2, 0), UNI, size=10.5)
    cell_text(edu.cell(2, 1), DEPT + "（學號 " + SID + "）", size=10.5)
    cell_text(edu.cell(2, 2), "112 學年度入學，預計民國 116 年 6 月畢業", size=10)
    set_widths(edu, [5.0, 6.4, 6.0])

    add_sub(doc, "實習／工作經歷")
    work = doc.add_table(rows=2, cols=4)
    work.style = "Table Grid"
    for i, h in enumerate(["公司名稱", "職稱", "工作內容", "工作起迄"]):
        cell_text(work.cell(0, i), h, size=10, bold=True, align="center", fill="E8EEF0")
    for i in range(4):
        cell_text(work.cell(1, i), "", size=10.5)
    set_widths(work, [4.2, 3.2, 6.0, 4.0])

    add_sub(doc, "社團經驗")
    club = doc.add_table(rows=5, cols=4)
    club.style = "Table Grid"
    for i, h in enumerate(["社團名稱", "職稱", "重要事蹟", "期間"]):
        cell_text(club.cell(0, i), h, size=10, bold=True, align="center", fill="E8EEF0")
    rows = [
        ("資管系學會", "活動部部長", "規劃活動和跨年級協調", "113 學年度"),
        ("資管三B", "校外賃居聯絡人", "", "114-1"),
        ("資管三B", "副班代", "", "114-2"),
        ("資管四B", "服務股長", "", "115-1"),
    ]
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell_text(club.cell(r, c), val, size=10.5)
    set_widths(club, [4.0, 3.4, 6.2, 3.8])

    add_sub(doc, "證照")
    cert = doc.add_table(rows=2, cols=3)
    cert.style = "Table Grid"
    for i, h in enumerate(["證照名稱", "發照單位", "發照日期"]):
        cell_text(cert.cell(0, i), h, size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(cert.cell(1, 0), "ERP 規劃師　GP0115533440", size=10.5)
    cell_text(cert.cell(1, 1), "中華企業資源規劃學會", size=10.5)
    cell_text(cert.cell(1, 2), "民國 114 年 12 月", size=10.5)
    set_widths(cert, [6.4, 6.4, 4.6])

    add_sub(doc, "競賽（志工／遊學／自學）")
    comp = doc.add_table(rows=5, cols=3)
    comp.style = "Table Grid"
    for i, h in enumerate(["獎第", "項目", "參賽期間"]):
        cell_text(comp.cell(0, i), h, size=10, bold=True, align="center", fill="E8EEF0")
    comps = [
        ("", "Kaggle Playground S6E9（公開分數 0.94540）", ""),
        ("", "Kaggriculture v2（Kaggle）：已提交並完成評分", ""),
        ("", "DrivenData 帕金森症挑戰（公開分數 0.5591，指標 log loss，越低越好）", ""),
        ("", "臺鐵數據力 2026（民國 115 年）數據資料創意發想競賽（學生組），作品 RailFlow-DSS 鐵道旅運負荷預測決策支援系統，繳交海選企劃", "2026/09"),
    ]
    for r, row in enumerate(comps, 1):
        for c, val in enumerate(row):
            cell_text(comp.cell(r, c), val, size=10)
    set_widths(comp, [2.4, 11.6, 3.4])

    add_sub(doc, "專長")
    skill = doc.add_table(rows=3, cols=2)
    skill.style = "Table Grid"
    cell_text(skill.cell(0, 0), "電腦技能", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(skill.cell(0, 1), "Python、pandas、scikit-learn、Streamlit、SQLite；GitHub Actions", size=10.5)
    cell_text(skill.cell(1, 0), "語文能力", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(skill.cell(1, 1), "", size=10.5)
    cell_text(skill.cell(2, 0), "其他", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(skill.cell(2, 1), "應用層存取控制與操作紀錄；資訊系統與資訊安全方向", size=10.5)
    set_widths(skill, [3.2, 14.2])

    doc.add_page_break()
    k2 = doc.add_paragraph()
    k2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    k2.paragraph_format.space_after = Pt(0)
    r2 = k2.add_run("靜宜大學 E-Portfolio")
    set_run_font(r2, size=10)
    add_heading_p(doc, "自傳")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    rs = sub.add_run(NAME + "　｜　" + UNI + DEPT + "四年級　｜　英文姓名 " + EN_NAME)
    set_run_font(rs, size=10)

    add_sub(doc, "成長背景")
    add_body(doc, "我民國 94 年 7 月 29 日出生。")

    add_sub(doc, "求學經歷")
    add_body(doc, "高中就讀楊子高中。大學就讀靜宜大學資訊管理學系，學號 411211325，112 學年度入學，四年級，預計民國 116 年 6 月畢業。")
    add_body(doc, "修過的課包括程式語言、C 語言程式設計、程式設計、機器學習、專案實作（一）和系統分析與設計。需求預測用機器學習課的模型。角色和補貨單據用系統分析與設計課的方法整理。")

    add_sub(doc, "工作經驗")
    add_body(doc, "113 學年度擔任資管系學會活動部部長，負責規劃活動和跨年級協調。114-1 擔任資管三B 校外賃居聯絡人。114-2 擔任資管三B 副班代。115-1 擔任資管四B 服務股長。")
    add_body(doc, "課外個人作品是 Retail Ops DSS，github.com/Miiduoa/retail-ops-dss 公開，固定版本 1535a052。系統分成檢視、門市營運、管理者三種角色。每個動作先查角色能不能做，權限表裡沒列的動作一律拒絕。門市營運人員只能看、只能改帳號授權範圍內的門市。超出帳號授權門市範圍的操作會被擋下並留下紀錄。數量很大的補貨核准、變更系統設定這類高影響操作，要近期再驗證一次身分才放行。成功和被拒絕的操作都寫進只能新增、不能修改或刪除的紀錄。")
    add_body(doc, "我寫了 24 項測試，以存取控制為主。固定版本 2026-10-06 在 GitHub Actions 上驗證通過。")

    add_sub(doc, "未來規劃")
    add_body(doc, "畢業後我想先做資訊系統開發或維護，再做帳號與存取權限管理、系統安全維護。方向是資訊與資訊安全。權限或程式一改，先用自動測試核對。")

    add_sub(doc, "證照、獎狀")
    att1 = doc.add_table(rows=2, cols=2)
    att1.style = "Table Grid"
    cell_text(att1.cell(0, 0), "名稱", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(att1.cell(0, 1), "期間", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(att1.cell(1, 0), "ERP 規劃師，中華企業資源規劃學會，GP0115533440", size=10.5)
    cell_text(att1.cell(1, 1), "民國 114 年 12 月", size=10.5)
    set_widths(att1, [12.4, 5.0])

    add_sub(doc, "實習、服務、社團證明")
    att2 = doc.add_table(rows=5, cols=2)
    att2.style = "Table Grid"
    cell_text(att2.cell(0, 0), "名稱", size=10, bold=True, align="center", fill="E8EEF0")
    cell_text(att2.cell(0, 1), "期間", size=10, bold=True, align="center", fill="E8EEF0")
    att_rows = [
        ("資管系學會活動部部長", "113 學年度"),
        ("資管三B 校外賃居聯絡人", "114-1"),
        ("資管三B 副班代", "114-2"),
        ("資管四B 服務股長", "115-1"),
    ]
    for r, row in enumerate(att_rows, 1):
        cell_text(att2.cell(r, 0), row[0], size=10.5)
        cell_text(att2.cell(r, 1), row[1], size=10.5)
    set_widths(att2, [12.4, 5.0])

    path = os.path.join(ROOT, DOCX_NAME)
    doc.save(path)
    return path


CSS = r"""
@font-face { font-family: "Noto Sans CJK TC"; font-weight: 400; src: url("file://""" + FONT_SANS + r""""); }
@font-face { font-family: "Noto Sans CJK TC"; font-weight: 700; src: url("file://""" + FONT_SANS + r""""); }
@font-face { font-family: "Noto Serif CJK TC"; font-weight: 400; src: url("file://""" + FONT_SERIF + r""""); }
html, body { background:#ffffff; color:#1c2733; font-family:"Noto Serif CJK TC", serif;
  font-size:10pt; line-height:1.42; margin:0; padding:0; word-break:keep-all; }
h1 { font-family:"Noto Sans CJK TC", sans-serif; font-weight:700; color:#1a4550;
  text-align:center; font-size:15pt; margin:2px 0 6px 0; }
h2 { font-family:"Noto Sans CJK TC", sans-serif; font-weight:700; color:#1a4550;
  font-size:11pt; border-bottom:1.1pt solid #1a4550; padding:0 0 2px 0; margin:7px 0 3px 0; }
.kicker { font-family:"Noto Sans CJK TC", sans-serif; font-size:8.5pt; color:#5e5a55; text-align:center; margin:0; }
.sub { font-family:"Noto Sans CJK TC", sans-serif; color:#5e5a55; text-align:center; font-size:9.5pt; margin:0 0 8px 0; }
p { margin:0 0 6px 0; text-align:justify; }
table { border-collapse:collapse; width:100%; margin:0 0 4px 0; font-size:9.6pt; }
th, td { border:0.6pt solid #1a4550; padding:3px 6px; vertical-align:middle; }
th, td.lbl { background:#e8eef0; font-family:"Noto Sans CJK TC", sans-serif; font-weight:700; color:#1c2733; white-space:nowrap; }
.photo { width:22mm; height:auto; display:block; margin:1.5mm auto; }
.keep { page-break-inside: avoid; }
.pagebreak { page-break-before:always; }
@page { size:A4; margin:13mm 14mm 14mm 14mm;
  @bottom-left { content:"顧晉瑋｜靜宜大學資訊管理學系"; font-family:"Noto Sans CJK TC"; font-size:8pt; color:#5e5a55; }
  @bottom-right { content:"第 " counter(page) " 頁"; font-family:"Noto Sans CJK TC"; font-size:8pt; color:#5e5a55; }
}
"""


def build_html():
    photo_uri = "file://" + PHOTO
    return f"""<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8"><style>{CSS}</style></head><body>
<p class="kicker">靜宜大學 E-Portfolio</p>
<h1>履歷表</h1>
<table>
<tr>
  <td class="lbl" style="width:16%">應徵職務</td><td colspan="3"></td>
  <td rowspan="7" style="width:20%;text-align:center;padding:4px;font-family:'Noto Sans CJK TC',sans-serif;font-size:8.5pt;color:#5e5a55">請貼照片<br>
    <img class="photo" src="{photo_uri}" alt="顧晉瑋">
  </td>
</tr>
<tr><td class="lbl">姓名</td><td>{NAME}</td><td class="lbl" style="width:12%">性別</td><td>{GENDER}</td></tr>
<tr><td class="lbl">出生年月日</td><td colspan="3">{BIRTH}</td></tr>
<tr><td class="lbl">服役狀況</td><td colspan="3">□免役　　□未役　　□待役　　□役畢　　年　　月</td></tr>
<tr><td class="lbl">聯絡電話</td><td colspan="3">{PHONE}</td></tr>
<tr><td class="lbl">電子信箱</td><td colspan="3">{EMAIL}</td></tr>
<tr><td class="lbl">聯絡地址</td><td colspan="3">{ADDR}</td></tr>
</table>
<h2>教育程度</h2>
<table>
<tr><th>學校</th><th>科系所</th><th>修業年限</th></tr>
<tr><td>{HS}</td><td></td><td></td></tr>
<tr><td>{UNI}</td><td>{DEPT}（學號 {SID}）</td><td>112 學年度入學，預計民國 116 年 6 月畢業</td></tr>
</table>
<h2>實習／工作經歷</h2>
<table>
<tr><th>公司名稱</th><th>職稱</th><th>工作內容</th><th>工作起迄</th></tr>
<tr><td></td><td></td><td></td><td></td></tr>
</table>
<h2>社團經驗</h2>
<table>
<tr><th>社團名稱</th><th>職稱</th><th>重要事蹟</th><th>期間</th></tr>
<tr><td>資管系學會</td><td>活動部部長</td><td>規劃活動和跨年級協調</td><td>113 學年度</td></tr>
<tr><td>資管三B</td><td>校外賃居聯絡人</td><td></td><td>114-1</td></tr>
<tr><td>資管三B</td><td>副班代</td><td></td><td>114-2</td></tr>
<tr><td>資管四B</td><td>服務股長</td><td></td><td>115-1</td></tr>
</table>
<h2>證照</h2>
<table>
<tr><th>證照名稱</th><th>發照單位</th><th>發照日期</th></tr>
<tr><td>ERP 規劃師　GP0115533440</td><td>中華企業資源規劃學會</td><td>民國 114 年 12 月</td></tr>
</table>
<div class="keep">
<h2>競賽（志工／遊學／自學）</h2>
<table>
<tr><th>獎第</th><th>項目</th><th>參賽期間</th></tr>
<tr><td></td><td>Kaggle Playground S6E9（公開分數 0.94540）</td><td></td></tr>
<tr><td></td><td>Kaggriculture v2（Kaggle）：已提交並完成評分</td><td></td></tr>
<tr><td></td><td>DrivenData 帕金森症挑戰（公開分數 0.5591，指標 log loss，越低越好）</td><td></td></tr>
<tr><td></td><td>臺鐵數據力 2026（民國 115 年）數據資料創意發想競賽（學生組），作品 RailFlow-DSS 鐵道旅運負荷預測決策支援系統，繳交海選企劃</td><td>2026/09</td></tr>
</table>
</div>
<div class="keep">
<h2>專長</h2>
<table>
<tr><td class="lbl">電腦技能</td><td>Python、pandas、scikit-learn、Streamlit、SQLite；GitHub Actions</td></tr>
<tr><td class="lbl">語文能力</td><td></td></tr>
<tr><td class="lbl">其他</td><td>應用層存取控制與操作紀錄；資訊系統與資訊安全方向</td></tr>
</table>
</div>
<div class="pagebreak"></div>
<p class="kicker">靜宜大學 E-Portfolio</p>
<h1>自傳</h1>
<p class="sub">{NAME}　｜　{UNI}{DEPT}四年級　｜　英文姓名 {EN_NAME}</p>
<h2>成長背景</h2>
<p>我民國 94 年 7 月 29 日出生。</p>
<h2>求學經歷</h2>
<p>高中就讀楊子高中。大學就讀靜宜大學資訊管理學系，學號 411211325，112 學年度入學，四年級，預計民國 116 年 6 月畢業。</p>
<p>修過的課包括程式語言、C 語言程式設計、程式設計、機器學習、專案實作（一）和系統分析與設計。需求預測用機器學習課的模型。角色和補貨單據用系統分析與設計課的方法整理。</p>
<h2>工作經驗</h2>
<p>113 學年度擔任資管系學會活動部部長，負責規劃活動和跨年級協調。114-1 擔任資管三B 校外賃居聯絡人。114-2 擔任資管三B 副班代。115-1 擔任資管四B 服務股長。</p>
<p>課外個人作品是 Retail Ops DSS，github.com/Miiduoa/retail-ops-dss 公開，固定版本 1535a052。系統分成檢視、門市營運、管理者三種角色。每個動作先查角色能不能做，權限表裡沒列的動作一律拒絕。門市營運人員只能看、只能改帳號授權範圍內的門市。超出帳號授權門市範圍的操作會被擋下並留下紀錄。數量很大的補貨核准、變更系統設定這類高影響操作，要近期再驗證一次身分才放行。成功和被拒絕的操作都寫進只能新增、不能修改或刪除的紀錄。</p>
<p>我寫了 24 項測試，以存取控制為主。固定版本 2026-10-06 在 GitHub Actions 上驗證通過。</p>
<h2>未來規劃</h2>
<p>畢業後我想先做資訊系統開發或維護，再做帳號與存取權限管理、系統安全維護。方向是資訊與資訊安全。權限或程式一改，先用自動測試核對。</p>
<h2>證照、獎狀</h2>
<table>
<tr><th>名稱</th><th>期間</th></tr>
<tr><td>ERP 規劃師，中華企業資源規劃學會，GP0115533440</td><td>民國 114 年 12 月</td></tr>
</table>
<h2>實習、服務、社團證明</h2>
<table>
<tr><th>名稱</th><th>期間</th></tr>
<tr><td>資管系學會活動部部長</td><td>113 學年度</td></tr>
<tr><td>資管三B 校外賃居聯絡人</td><td>114-1</td></tr>
<tr><td>資管三B 副班代</td><td>114-2</td></tr>
<tr><td>資管四B 服務股長</td><td>115-1</td></tr>
</table>
</body></html>"""


def strip_meta(doc):
    doc.set_metadata({
        "title": "", "author": "", "subject": "",
        "keywords": "", "creator": "", "producer": "",
        "creationDate": "", "modDate": "",
    })
    doc.del_xml_metadata()
    for page in doc:
        for link in page.get_links():
            try:
                page.delete_link(link)
            except Exception:
                pass


def save_pdf(html):
    html_path = os.path.join(ROOT, "resume.html")
    open(html_path, "w", encoding="utf-8").write(html)
    raw = os.path.join(ROOT, "resume.raw.pdf")
    HTML(filename=html_path, base_url=ROOT).write_pdf(raw)
    d = pymupdf.open(raw)
    strip_meta(d)
    dest = os.path.join(ROOT, PDF_NAME)
    tmp = dest + ".tmp"
    d.save(tmp, garbage=4, deflate=True, no_new_id=True)
    bb = open(tmp, "rb").read()
    bb = re.sub(rb"% Written by MuPDF[^\r\n]*", lambda m: b"%" + b" " * (len(m.group(0)) - 1), bb)
    open(dest, "wb").write(bb)
    os.remove(tmp)
    return dest, d.page_count


def copy_out(src, dest):
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    shutil.copy2(src, dest)
    print("copied", dest, os.path.getsize(dest))


def main():
    html = build_html()
    joined = html
    for bad in ["而是", "而不是", "而不只", "活動長", "學業節錄", "73.74",
                "目前只做到", "AI 工具", "碩士班", "報考", "免責"]:
        assert bad not in joined, bad
    assert "活動部部長" in joined
    assert "GP0115533440" in joined
    assert "已提交並完成評分" in joined
    assert "臺鐵數據力 2026" in joined
    assert "1535a052" in joined
    assert "0966198826" in joined
    assert "學府路200號14樓之1" in joined
    assert "楊子高中" in joined
    assert "程式設計 94" not in joined
    assert "程式設計 81" not in joined

    docx_path = build_docx()
    pdf_path, n = save_pdf(html)
    d = pymupdf.open(pdf_path)
    text = "\n".join(p.get_text("text") for p in d)
    meta = d.metadata
    print("pages", d.page_count, "pdf_bytes", os.path.getsize(pdf_path),
          "docx_bytes", os.path.getsize(docx_path),
          "creator", repr(meta.get("creator")))
    assert 2 <= d.page_count <= 3
    assert os.path.getsize(pdf_path) <= 2 * 1024 * 1024
    assert os.path.getsize(docx_path) <= 2 * 1024 * 1024
    assert not meta.get("creator")
    assert "活動部部長" in text
    assert "而是" not in text
    assert "73.74" not in text
    assert "碩士班" not in text
    pdir = os.path.join(OUT_DIR, "pages")
    for dest_pages in (pdir, os.path.join(STORE, "pages"), os.path.join(WS_DIR, "pages")):
        os.makedirs(dest_pages, exist_ok=True)
        for fn in os.listdir(dest_pages):
            if fn.endswith(".png"):
                os.remove(os.path.join(dest_pages, fn))
    for i, p in enumerate(d):
        pix = p.get_pixmap(matrix=pymupdf.Matrix(1.55, 1.55), alpha=False)
        pix.save(os.path.join(pdir, f"p{i+1:02d}.png"))
    for dest_root in (OUT_DIR, STORE, WS_DIR):
        copy_out(docx_path, os.path.join(dest_root, DOCX_NAME))
        copy_out(pdf_path, os.path.join(dest_root, PDF_NAME))
        copy_out(docx_path, os.path.join(dest_root, "pu_resume_r1.docx"))
        copy_out(pdf_path, os.path.join(dest_root, "pu_resume_r1.pdf"))
        copy_out(__file__, os.path.join(dest_root, "gen_pu_resume_r1.py"))
    for fn in os.listdir(pdir):
        copy_out(os.path.join(pdir, fn), os.path.join(STORE, "pages", fn))
        copy_out(os.path.join(pdir, fn), os.path.join(WS_DIR, "pages", fn))
    h_pdf = hashlib.sha256(open(pdf_path, "rb").read()).hexdigest()
    h_doc = hashlib.sha256(open(docx_path, "rb").read()).hexdigest()
    sums = f"{h_pdf}  {PDF_NAME}\n{h_doc}  {DOCX_NAME}\n"
    for dest_root in (OUT_DIR, STORE, WS_DIR):
        open(os.path.join(dest_root, "SHA256SUMS.txt"), "w").write(sums)
    print(sums)
    print("done pages", n)


if __name__ == "__main__":
    main()
