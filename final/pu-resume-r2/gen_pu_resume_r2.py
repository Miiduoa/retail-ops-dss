# -*- coding: utf-8 -*-
"""靜宜 E-Portfolio 履歷自傳 r2：官方 autobiography_2020.doc 為底，只填格子。"""
import os, re, shutil, hashlib, subprocess
from copy import deepcopy

from docx import Document
from docx.shared import Cm, Pt, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf

ROOT = "/tmp/pu-resume-r2"
TPL_DOCX = "/tmp/pu-tpl/autobiography_2020.docx"
PHOTO = "/tmp/pu-resume-r1/photo.jpg"
OUT_DIR = "/opt/cursor/artifacts/pu-resume-r2"
STORE = "/cursor/stores/self/final/pu-resume-r2"
WS_DIR = "/workspace/final/pu-resume-r2"
DOCX_NAME = "履歷自傳_顧晉瑋.docx"
PDF_NAME = "履歷自傳_顧晉瑋.pdf"

NAME = "顧晉瑋"
GENDER = "男"
BIRTH = "民國 94 年 7 月 29 日"
PHONE = "0966198826"
EMAIL = "demohan513@gmail.com"
ADDR = "402 台中市南區學府路200號14樓之1"
JOB = "資安工程師、軟體工程師"
SID = "411211325"


def set_run_font(run, size=13, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "標楷體")
    rFonts.set(qn("w:cs"), "Times New Roman")


def tcs(row):
    return row._tr.findall(qn("w:tc"))


def cell(table, ri, ti):
    return _Cell(tcs(table.rows[ri])[ti], table)


def write(table, ri, ti, text, size=11):
    c = cell(table, ri, ti)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size)
    return c


def fill_name_period(table, ri, ti, name, period):
    c = cell(table, ri, ti)
    paras = c.paragraphs
    # keep 「名稱：」「期間：」 labels, append facts
    if len(paras) >= 2:
        paras[0].text = ""
        r = paras[0].add_run("名稱：" + name)
        set_run_font(r, size=12)
        paras[1].text = ""
        r = paras[1].add_run("期間：" + period)
        set_run_font(r, size=12)
    else:
        write(table, ri, ti, "名稱：" + name + "\n期間：" + period, size=12)


def add_photo(table):
    c = cell(table, 1, 4)
    # keep 請貼照片, add image below
    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(PHOTO, width=Cm(2.4))


def fill_para(p, text, size=13):
    p.text = ""
    run = p.add_run(text)
    set_run_font(run, size=size)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.4


def insert_after(p, text, size=13):
    new_el = deepcopy(p._p)
    # clear runs
    for child in list(new_el):
        if child.tag == qn("w:r"):
            new_el.remove(child)
    p._p.addnext(new_el)
    np = Paragraph(new_el, p._parent)
    fill_para(np, text, size=size)
    return np


def fill_docx():
    shutil.copy2(TPL_DOCX, os.path.join(ROOT, "base.docx"))
    doc = Document(os.path.join(ROOT, "base.docx"))
    t = doc.tables[2]

    write(t, 0, 1, JOB)
    write(t, 1, 1, NAME)
    write(t, 1, 3, GENDER)
    write(t, 2, 1, BIRTH)
    write(t, 4, 1, PHONE)
    write(t, 5, 1, EMAIL)
    write(t, 6, 1, ADDR)
    add_photo(t)

    write(t, 9, 0, "楊子高中")
    # 科系所、修業年限 unknown
    write(t, 10, 0, "靜宜大學")
    write(t, 10, 1, "資訊管理學系（學號 " + SID + "）")
    write(t, 10, 2, "112入學，預計116年6月畢業", size=11)

    # 實習/工作經歷：無，留空

    write(t, 18, 0, "資管系學會")
    write(t, 18, 1, "活動部部長")
    write(t, 18, 2, "規劃活動和跨年級協調")
    write(t, 18, 3, "113 學年度")
    write(t, 19, 0, "資管三B")
    write(t, 19, 1, "校外賃居聯絡人")
    write(t, 19, 3, "114-1")
    write(t, 20, 0, "資管三B、四B")
    write(t, 20, 1, "副班代／服務股長")
    write(t, 20, 3, "114-2、115-1")

    write(t, 23, 0, "ERP 規劃師　GP0115533440", size=11)
    write(t, 23, 1, "中華企業資源規劃學會")
    write(t, 23, 2, "民國 114 年 12 月")

    write(t, 27, 1, "Kaggle Playground S6E9（公開分數 0.94540）；Kaggriculture v2（Kaggle）：已提交並完成評分", size=11)
    write(t, 28, 1, "DrivenData 帕金森症挑戰（公開分數 0.5591，指標 log loss，越低越好）；臺鐵數據力 2026（民國 115 年）數據資料創意發想競賽（學生組），作品 RailFlow-DSS，2026/09 繳交海選企劃", size=11)

    write(t, 30, 1, "Python、pandas、scikit-learn、Streamlit、SQLite；GitHub Actions")
    write(t, 32, 1, "應用層存取控制與操作紀錄；資訊系統與資訊安全方向")

    def after_heading(title):
        paras = doc.paragraphs
        for i, p in enumerate(paras):
            if p.text.replace(" ", "") == title:
                return paras[i + 1]
        raise KeyError(title)

    p_grow = after_heading("成長背景")
    p_school = after_heading("求學經歷")
    p_work = after_heading("工作經驗")
    p_plan = after_heading("未來規劃")
    fill_para(p_grow, "我民國 94 年 7 月 29 日出生。")
    fill_para(p_school, "高中就讀楊子高中。大學就讀靜宜大學資訊管理學系，學號 411211325，112 學年度入學，四年級，預計民國 116 年 6 月畢業。")
    insert_after(
        p_school,
        "修過的課包括程式語言、C 語言程式設計、程式設計、機器學習、專案實作（一）和系統分析與設計。需求預測用機器學習課的模型。角色和補貨單據用系統分析與設計課的方法整理。",
    )
    fill_para(p_work, "113 學年度擔任資管系學會活動部部長，負責規劃活動和跨年級協調。114-1 擔任資管三B 校外賃居聯絡人。114-2 擔任資管三B 副班代。115-1 擔任資管四B 服務股長。")
    p_work_b = insert_after(
        p_work,
        "課外個人作品是 Retail Ops DSS，github.com/Miiduoa/retail-ops-dss 公開，固定版本 1535a052。系統分成檢視、門市營運、管理者三種角色。每個動作先查角色能不能做，權限表裡沒列的動作一律拒絕。門市營運人員只能看、只能改帳號授權範圍內的門市。超出帳號授權門市範圍的操作會被擋下並留下紀錄。數量很大的補貨核准、變更系統設定這類高影響操作，要近期再驗證一次身分才放行。成功和被拒絕的操作都寫進只能新增、不能修改或刪除的紀錄。",
    )
    insert_after(
        p_work_b,
        "我寫了 24 項測試，以存取控制為主。固定版本 2026-10-06 在 GitHub Actions 上驗證通過。",
    )
    fill_para(p_plan, "畢業後我想先做資訊系統開發或維護，再做帳號與存取權限管理、系統安全維護。方向是資訊與資訊安全。權限或程式一改，先用自動測試核對。")

    for p in doc.paragraphs:
        if "證照、獎狀" in p.text:
            p.paragraph_format.keep_with_next = True

    # 證照、獎狀
    att1 = doc.tables[3]
    fill_name_period(att1, 1, 0, "ERP規劃師 GP0115533440", "民國114年12月")

    # 實習、服務、社團證明
    att2 = doc.tables[4]
    fill_name_period(att2, 1, 0, "資管系學會活動部部長", "113 學年度")
    fill_name_period(att2, 1, 1, "資管三B 校外賃居聯絡人", "114-1")
    fill_name_period(att2, 3, 0, "資管三B 副班代", "114-2")
    fill_name_period(att2, 3, 1, "資管四B 服務股長", "115-1")

    dest = os.path.join(ROOT, DOCX_NAME)
    doc.save(dest)
    return dest


def strip_meta(doc):
    doc.set_metadata({
        "title": "", "author": "", "subject": "",
        "keywords": "", "creator": "", "producer": "",
        "creationDate": "", "modDate": "",
    })
    doc.del_xml_metadata()
    for page in doc:
        for link in page.get_links():
            try:
                page.delete_link(link)
            except Exception:
                pass


def to_pdf(docx_path):
    raw = os.path.join(ROOT, "resume.raw.pdf")
    subprocess.check_call([
        "soffice", "--headless", "--convert-to", "pdf:writer_pdf_Export",
        "--outdir", ROOT, docx_path,
    ], stdout=subprocess.DEVNULL, stderr=subprocess.STDOUT)
    produced = os.path.join(ROOT, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    shutil.move(produced, raw)
    d = pymupdf.open(raw)
    strip_meta(d)
    dest = os.path.join(ROOT, PDF_NAME)
    tmp = dest + ".tmp"
    d.save(tmp, garbage=4, deflate=True, no_new_id=True)
    bb = open(tmp, "rb").read()
    bb = re.sub(rb"% Written by MuPDF[^\r\n]*", lambda m: b"%" + b" " * (len(m.group(0)) - 1), bb)
    open(dest, "wb").write(bb)
    os.remove(tmp)
    return dest, d.page_count


def copy_out(src, dest):
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    try:
        shutil.copy2(src, dest)
    except OSError as e:
        print("copy fail", dest, e)


def main():
    for d in (ROOT, OUT_DIR, STORE, WS_DIR,
              os.path.join(OUT_DIR, "pages"),
              os.path.join(STORE, "pages"),
              os.path.join(WS_DIR, "pages"),
              os.path.join(STORE, "tpl")):
        os.makedirs(d, exist_ok=True)

    docx_path = fill_docx()
    pdf_path, n = to_pdf(docx_path)
    d = pymupdf.open(pdf_path)
    text = "\n".join(p.get_text("text") for p in d)
    meta = d.metadata
    print("pages", d.page_count, "pdf", os.path.getsize(pdf_path),
          "docx", os.path.getsize(docx_path), "creator", repr(meta.get("creator")))

    joined = text
    for bad in ["而是", "而不是", "而不只", "活動長", "學業節錄", "73.74",
                "目前只做到", "AI 工具", "碩士班", "報考", "免責",
                "程式設計 94", "程式設計 81"]:
        assert bad not in joined, bad
    for need in ["活動部部長", "GP0115533440", "已提交並完成評分", "臺鐵數據力 2026",
                 "1535a052", "0966198826", "學府路200", "楊子高中",
                 "資安工程師", "軟體工程師", "411211325"]:
        assert need in joined, need
    assert 5 <= d.page_count <= 7
    assert os.path.getsize(pdf_path) <= 2 * 1024 * 1024
    assert os.path.getsize(docx_path) <= 2 * 1024 * 1024
    assert not meta.get("creator")

    pdir = os.path.join(ROOT, "pages")
    os.makedirs(pdir, exist_ok=True)
    for fn in os.listdir(pdir):
        if fn.endswith(".png"):
            os.remove(os.path.join(pdir, fn))
    for dest_pages in (os.path.join(STORE, "pages"), os.path.join(WS_DIR, "pages"), os.path.join(OUT_DIR, "pages")):
        os.makedirs(dest_pages, exist_ok=True)
        for fn in os.listdir(dest_pages):
            if fn.endswith(".png"):
                try:
                    os.remove(os.path.join(dest_pages, fn))
                except OSError:
                    pass
    for i, p in enumerate(d):
        pix = p.get_pixmap(matrix=pymupdf.Matrix(1.55, 1.55), alpha=False)
        pix.save(os.path.join(pdir, f"p{i+1:02d}.png"))

    for dest_root in (STORE, WS_DIR, OUT_DIR):
        copy_out(docx_path, os.path.join(dest_root, DOCX_NAME))
        copy_out(pdf_path, os.path.join(dest_root, PDF_NAME))
        copy_out(docx_path, os.path.join(dest_root, "pu_resume_r2.docx"))
        copy_out(pdf_path, os.path.join(dest_root, "pu_resume_r2.pdf"))
        copy_out(__file__, os.path.join(dest_root, "gen_pu_resume_r2.py"))
    for fn in os.listdir(pdir):
        copy_out(os.path.join(pdir, fn), os.path.join(STORE, "pages", fn))
        copy_out(os.path.join(pdir, fn), os.path.join(WS_DIR, "pages", fn))
    for fn in os.listdir("/tmp/pu-tpl/pages"):
        copy_out(os.path.join("/tmp/pu-tpl/pages", fn), os.path.join(STORE, "tpl", fn))

    h_pdf = hashlib.sha256(open(pdf_path, "rb").read()).hexdigest()
    h_doc = hashlib.sha256(open(docx_path, "rb").read()).hexdigest()
    sums = f"{h_pdf}  {PDF_NAME}\n{h_doc}  {DOCX_NAME}\n"
    for dest_root in (OUT_DIR, STORE, WS_DIR):
        open(os.path.join(dest_root, "SHA256SUMS.txt"), "w").write(sums)
    print(sums)
    print("done pages", n)


if __name__ == "__main__":
    main()
